"""Reading uploaded (untrusted) .xlsx files for the personnel and user import screens.

`openpyxl.load_workbook` parses every part of an uploaded workbook (zip container, XML parts,
embedded images through Pillow). A corrupt or hostile file therefore raises whatever the zip,
XML, openpyxl or Pillow layer raises. This module turns exactly those parse/format failures into
one domain error with a fixed Turkish message, before any database work starts. The original
exception chain is dropped (``from None``) so no parser text, archive member name or file content
reaches a handler log; only the exception class name is logged.
"""
from __future__ import annotations

import logging
import lzma
import zipfile
import zlib
from typing import IO, Any
from xml.etree.ElementTree import ParseError

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from PIL.Image import DecompressionBombError

logger = logging.getLogger(__name__)

EXCEL_WORKBOOK_UNREADABLE_MESSAGE = (
    "Excel dosyası okunamadı. Dosya bozuk, şifreli veya desteklenmeyen içerik barındırıyor olabilir; "
    "dosyayı Excel'de açıp .xlsx olarak yeniden kaydedin ve tekrar yükleyin. "
    "Hiçbir kayıt veritabanına yazılmadı."
)

# What openpyxl 3.1.5 load_workbook(data_only=True) and Worksheet.iter_rows raised for corrupt and
# hostile workbooks under Pillow 12.3.0 and 13.0.0 (EOFError: stdlib zipfile, short member). The try
# blocks below contain only those two parser calls, so the broad built-in types here can only come
# from parsing the uploaded bytes, not from BYS360 code. OSError also covers an unreadable upload
# temp file and RuntimeError a RecursionError in a parser; both are reported as an unreadable file.
# AttributeError, NameError, ImportError and MemoryError are deliberately not caught.
_UNREADABLE_WORKBOOK_ERRORS: tuple[type[BaseException], ...] = (
    zipfile.BadZipFile,  # not a zip, truncated, OLE/CFB (encrypted .xlsx), bad CRC-32
    zlib.error,  # corrupt deflate member
    lzma.LZMAError,  # corrupt LZMA member
    EOFError,  # member shorter than declared
    OSError,  # "File contains no valid workbook part"; corrupt bzip2 member
    RuntimeError,  # encrypted zip member; NotImplementedError: unsupported compression
    ParseError,  # malformed XML part (SyntaxError subclass, not wrapped by openpyxl)
    ValueError,  # openpyxl read() wrapper (e.g. Pillow 13 truncated PNG chunk), row bound
    KeyError,  # missing part or dangling relationship id
    IndexError,  # shared-string or style index out of range
    TypeError,  # openpyxl descriptor type validation
    DecompressionBombError,  # embedded image over 2 x PIL.Image.MAX_IMAGE_PIXELS
)


class ExcelWorkbookReadError(ValueError):
    """Carries only the fixed user message; never parser text or file content."""

    def __init__(self, message: str = EXCEL_WORKBOOK_UNREADABLE_MESSAGE) -> None:
        super().__init__(message)
        self.user_message = message


def _log_unreadable(code: str, kind: type[Any]) -> None:
    logger.warning(
        "BYS360 Excel çalışma kitabı okunamadı | code=%s | exc_type=%s.%s",
        code,
        kind.__module__,
        kind.__qualname__,
    )


def read_active_sheet_rows(file_storage: IO[bytes] | Any) -> list[tuple[Any, ...]]:
    """Return every row of the active worksheet as value tuples, or raise ExcelWorkbookReadError."""
    try:
        workbook = load_workbook(file_storage, data_only=True)
    except _UNREADABLE_WORKBOOK_ERRORS as exc:
        _log_unreadable("EXCEL_WORKBOOK_UNREADABLE", type(exc))
        raise ExcelWorkbookReadError() from None
    worksheet = workbook.active
    if not isinstance(worksheet, Worksheet):
        # no sheet, activeTab out of range, missing sheet part or a chartsheet
        _log_unreadable("EXCEL_WORKBOOK_NO_WORKSHEET", type(worksheet))
        raise ExcelWorkbookReadError()
    try:
        return list(worksheet.iter_rows(values_only=True))
    except _UNREADABLE_WORKBOOK_ERRORS as exc:
        _log_unreadable("EXCEL_WORKBOOK_UNREADABLE", type(exc))
        raise ExcelWorkbookReadError() from None
