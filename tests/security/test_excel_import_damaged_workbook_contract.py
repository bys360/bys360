"""Damaged or hostile Excel uploads end in a controlled refusal, never a 500 or a partial write.

Both personnel Excel routes read the upload with ``openpyxl.load_workbook``, which hands every
embedded picture to Pillow. Measured against openpyxl 3.1.5 with Pillow 12.3.0 and 13.0.0, a
damaged ``.xlsx`` raises ``zipfile.BadZipFile``, ``zlib.error``, ``KeyError``,
``xml.etree.ElementTree.ParseError``, ``AttributeError``, ``ValueError`` or
``NotImplementedError``, and a picture declaring more pixels than Pillow's limit raises
``PIL.Image.DecompressionBombError`` -- openpyxl only swallows ``OSError`` for unreadable
pictures, so the bomb error reaches the route.

- ``/personnel/excel-upload`` already turned these into a flash message and a rollback.
- ``/admin/users/import`` read the workbook outside any handler: every case above was a 500.

Each test goes through the real HTTP route with a real admin web session on a throwaway SQLite
file, and checks the database afterwards: no user row and no ``account.create`` audit row.
A valid workbook with a real PNG inside is imported by both routes (Pillow reads it), and a
failure in the middle of the rows rolls back the rows that were already flushed.
"""
from __future__ import annotations

import io
import struct
import tempfile
import uuid
import zipfile
import zlib
from pathlib import Path
from types import SimpleNamespace

import pytest

PASSWORD = "ExcelDamagedWorkbookTest1!"
ADMIN_SICIL = "XD01ADM"
_DB_ROOT = Path(tempfile.gettempdir()) / "bys360" / "excel_damaged_workbook_contract"
HEADERS = ("Sicil No", "Ad", "Soyad", "E-posta", "Unvan", "Birim", "Üst Birim", "Rol")
IMPORT_PATHS = ("/personnel/excel-upload", "/admin/users/import")


def _make_app(monkeypatch, uri):
    for key, value in {
        "APP_ENV": "testing",
        "FLASK_ENV": "testing",
        "SECRET_KEY": "test-secret-key-excel-damaged-workbook",
        "DATABASE_URL": uri,
        "FLASK_SKIP_SCHEMA_VALIDATION": "1",
        "AUTO_REPAIR_SCHEMA": "false",
        "STRICT_SCHEMA_CHECK": "false",
        "REQUIRE_DOTENV_FILE": "false",
        "STRICT_ENV_VALIDATION": "false",
        "WTF_CSRF_ENABLED": "false",
        "SCHEDULER_ENABLED": "false",
        "MAIL_SUPPRESS_SEND": "true",
        "LOGIN_FORCE_CAPTCHA_FOR_UNKNOWN_USER": "false",
        "DEFAULT_FIRST_LOGIN_PASSWORD": "excel-damaged-test-first-login",
    }.items():
        monkeypatch.setenv(key, value)
    from app import create_app
    from config import Config

    monkeypatch.setattr(Config, "APP_ENV", "testing")
    monkeypatch.setattr(Config, "SQLALCHEMY_DATABASE_URI", uri)
    monkeypatch.setattr(Config, "SQLALCHEMY_ENGINE_OPTIONS", {})
    flask_app = create_app()
    flask_app.config.update(TESTING=True, WTF_CSRF_ENABLED=False, SQLALCHEMY_DATABASE_URI=uri)
    return flask_app


@pytest.fixture
def env(monkeypatch, install_interim_notes_schema):
    _DB_ROOT.mkdir(parents=True, exist_ok=True)
    db_file = _DB_ROOT / f"xd_{uuid.uuid4().hex}.sqlite3"
    app = _make_app(monkeypatch, "sqlite:///" + db_file.as_posix())
    from app.extensions import db
    from app.models import User
    from app.services import runtime_schema

    with app.app_context():
        db.create_all()
        runtime_schema.provision_all()
        install_interim_notes_schema(db.engine)
        admin = User(
            sicil_no=ADMIN_SICIL,
            email="xd01adm@example.gov.tr",
            ad="Kisi",
            soyad="Admin",
            unvan="Uzman",
            role="admin",
            birim="XD Birim",
            ust_birim="XD Ust Birim",
            is_active=True,
            must_change_password=False,
            must_set_security_question=False,
            is_first_login=False,
        )
        admin.set_password(PASSWORD)
        db.session.add(admin)
        db.session.commit()
    yield SimpleNamespace(app=app)
    with app.app_context():
        db.session.remove()
        db.engine.dispose()
    db_file.unlink(missing_ok=True)


def _admin_client(env):
    client = env.app.test_client()
    response = client.post("/login", data={"sicil_or_email": ADMIN_SICIL, "password": PASSWORD})
    assert response.status_code == 302 and "/login" not in response.headers.get("Location", "")
    return client


def _snapshot(env):
    """(user sicils, number of account.create audit rows)."""
    from app.models import AuditLog, User

    with env.app.app_context():
        sicils = sorted(user.sicil_no for user in User.query.all())
        creates = AuditLog.query.filter_by(action="account.create").count()
    return sicils, creates


# --------------------------------------------------------------------------------------
# Workbooks
# --------------------------------------------------------------------------------------


def _png(width=8, height=8) -> bytes:
    from PIL import Image

    buffer = io.BytesIO()
    Image.new("RGB", (width, height), (200, 30, 30)).save(buffer, "PNG")
    return buffer.getvalue()


def _bomb_png(width=30000, height=30000) -> bytes:
    """A well-formed PNG that declares width*height pixels (over Pillow's bomb limit) in ~1 KB."""

    def chunk(kind: bytes, data: bytes) -> bytes:
        return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data) & 0xFFFFFFFF)

    header = struct.pack(">IIBBBBB", width, height, 1, 0, 0, 0, 0)
    body = zlib.compress(b"\x00" * (width // 8 + 1) * 10)
    return b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", header) + chunk(b"IDAT", body) + chunk(b"IEND", b"")


def _row(sicil):
    return (sicil, "Kisi", sicil, f"{sicil.lower()}@example.gov.tr", "Uzman", "XD Birim", "XD Ust Birim", "personel")


def _workbook(rows, *, picture: bytes | None = None) -> bytes:
    """A real .xlsx; with ``picture`` its embedded image part carries exactly those bytes."""
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as XLImage

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(list(HEADERS))
    for row in rows:
        sheet.append(list(row))
    if picture is not None:
        sheet.add_image(XLImage(io.BytesIO(_png())), "K2")
    buffer = io.BytesIO()
    workbook.save(buffer)
    if picture is None:
        return buffer.getvalue()
    return _replace_parts(buffer.getvalue(), lambda name, data: picture if name.startswith("xl/media/") else data)


def _replace_parts(xlsx: bytes, change) -> bytes:
    """Rewrite each zip part through ``change(name, data)``; ``None`` drops the part."""
    source = zipfile.ZipFile(io.BytesIO(xlsx))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            data = change(info.filename, source.read(info.filename))
            if data is not None:
                target.writestr(info.filename, data)
    return out.getvalue()


def _bit_flipped(xlsx: bytes) -> bytes:
    """Corrupt the compressed body of the first sheet (its local header stays intact)."""
    data = bytearray(xlsx)
    entry = zipfile.ZipFile(io.BytesIO(xlsx)).getinfo("xl/worksheets/sheet1.xml")
    name_len, extra_len = struct.unpack("<HH", data[entry.header_offset + 26: entry.header_offset + 30])
    start = entry.header_offset + 30 + name_len + extra_len
    for offset in range(start + 2, start + min(entry.compress_size, 40)):
        data[offset] ^= 0xA5
    return bytes(data)


def _damaged_workbooks() -> dict[str, bytes]:
    good = _workbook([_row("XDNEW01")])
    return {
        "not_a_zip": b"PK\x03\x04 this is not a workbook",
        "empty_file": b"",
        "truncated_zip": good[: len(good) // 2],
        "corrupt_sheet_stream": _bit_flipped(good),
        "missing_workbook_part": _replace_parts(good, lambda name, data: None if name == "xl/workbook.xml" else data),
        "garbled_sheet_xml": _replace_parts(
            good, lambda name, data: b"<worksheet><not-closed" if name == "xl/worksheets/sheet1.xml" else data
        ),
        "decompression_bomb_picture": _workbook([_row("XDNEW01")], picture=_bomb_png()),
    }


DAMAGED = _damaged_workbooks()


def _post(client, path, payload: bytes, **form):
    data = {"excel_file": (io.BytesIO(payload), "personel.xlsx"), **form}
    return client.post(path, data=data, content_type="multipart/form-data")


# --------------------------------------------------------------------------------------
# The damaged workbooks really fail inside openpyxl (no false green from a lenient reader)
# --------------------------------------------------------------------------------------


@pytest.mark.parametrize("case", sorted(DAMAGED))
def test_each_damaged_workbook_really_fails_to_load(case):
    from openpyxl import load_workbook

    with pytest.raises(Exception):  # noqa: B017 - which error is the point of the route tests
        workbook = load_workbook(io.BytesIO(DAMAGED[case]), data_only=True)
        list(workbook.active.iter_rows(values_only=True))


def test_the_bomb_picture_reaches_the_caller_as_decompression_bomb_error():
    from openpyxl import load_workbook
    from PIL import Image

    with pytest.raises(Image.DecompressionBombError):
        load_workbook(io.BytesIO(DAMAGED["decompression_bomb_picture"]), data_only=True)


# --------------------------------------------------------------------------------------
# Routes: controlled refusal, nothing written
# --------------------------------------------------------------------------------------


@pytest.mark.parametrize("case", sorted(DAMAGED))
def test_admin_user_import_refuses_a_damaged_workbook_without_a_500_or_a_write(env, case):
    client = _admin_client(env)
    before = _snapshot(env)

    response = _post(client, "/admin/users/import", DAMAGED[case])

    assert response.status_code == 200, (case, response.status_code)
    body = response.get_data(as_text=True)
    assert "Excel dosyası okunamadı" in body
    assert "Hiçbir kayıt veritabanına yazılmadı" in body
    assert _snapshot(env) == before


@pytest.mark.parametrize("case", sorted(DAMAGED))
def test_personnel_excel_upload_refuses_a_damaged_workbook_without_a_500_or_a_write(env, case):
    client = _admin_client(env)
    before = _snapshot(env)

    response = _post(client, "/personnel/excel-upload", DAMAGED[case])

    assert response.status_code in (200, 302), (case, response.status_code)
    assert _snapshot(env) == before


# --------------------------------------------------------------------------------------
# A real PNG inside a valid workbook is read and the rows are imported
# --------------------------------------------------------------------------------------


@pytest.mark.parametrize("path", IMPORT_PATHS)
def test_valid_workbook_with_an_embedded_png_is_imported(env, path):
    client = _admin_client(env)
    payload = _workbook([_row("XDPNG01"), _row("XDPNG02")], picture=_png(32, 32))

    response = _post(client, path, payload)

    assert response.status_code in (200, 302)
    sicils, creates = _snapshot(env)
    assert {"XDPNG01", "XDPNG02"} <= set(sicils)
    assert creates == 2


# --------------------------------------------------------------------------------------
# A failure in the middle of the rows rolls back the rows already flushed
# --------------------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("path", "module_name"),
    (("/personnel/excel-upload", "app.admin.routes"), ("/admin/users/import", "app.admin.ops_import_services")),
)
def test_a_failure_on_a_later_row_writes_none_of_the_rows(env, monkeypatch, path, module_name):
    import importlib

    module = importlib.import_module(module_name)
    real_assign = module.assign_user_performance_category
    calls = {"count": 0}

    def fail_on_third_row(user, *args, **kwargs):
        calls["count"] += 1
        if calls["count"] == 3:
            raise RuntimeError("synthetic failure on the third row")
        return real_assign(user, *args, **kwargs)

    monkeypatch.setattr(module, "assign_user_performance_category", fail_on_third_row)
    client = _admin_client(env)
    before = _snapshot(env)

    response = _post(client, path, _workbook([_row("XDRB01"), _row("XDRB02"), _row("XDRB03")]))

    assert calls["count"] == 3  # the first two rows really went through before the failure
    assert response.status_code in (200, 302)
    assert _snapshot(env) == before
