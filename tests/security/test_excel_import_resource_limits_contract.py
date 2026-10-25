"""Issue #104 contract: the three Excel import routes refuse resource-exhausting workbooks.

Routes: /admin/users/import, /personnel/excel-upload, /performance/import/history.
Before the fix each of them called ``load_workbook(data_only=True)`` + ``list(ws.iter_rows())``
on the upload: a 5 KB file with one cell at XFD1048576 cost ~143 s and 5.4 GB RSS (route-level
measurement in #104), a bold font on H1048576 42-83 s. Now every route reads through
``app.services.excel_import_limits.read_import_rows``:

- hostile or damaged files get a fixed Turkish message, never a 500;
- nothing is written: users, units, audit rows and import batches stay as they were, including
  when valid rows come before the offending cell;
- valid files still import;
- authorization is unchanged: a non-admin never reaches the reader.

Every request goes through the real route with a web session from ``POST /login`` (R05 fixture).
"""
from __future__ import annotations

import io
import zipfile
from datetime import date

import pytest
from openpyxl import Workbook
from openpyxl.styles import Font

from app.services.excel_import_limits import MESSAGES
from tests.security import test_r05_account_change_authority_contract as _r05
from tests.security.test_r05_account_change_authority_contract import (
    SICIL,
    _user_by_sicil,
    _web,
    _xlsx_row,
)

env = _r05.env  # the R05 fixture: seeded accounts, temp SQLite

HEADERS = ("Sicil No", "Ad", "Soyad", "E-posta", "Unvan", "Birim", "Üst Birim", "Rol")
ADMIN_IMPORT = "/admin/users/import"
PERSONNEL_UPLOAD = "/personnel/excel-upload"
HISTORY_IMPORT = "/performance/import/history"
ROUTES = (ADMIN_IMPORT, PERSONNEL_UPLOAD, HISTORY_IMPORT)
ROUTE_MODULES = {
    ADMIN_IMPORT: "app.admin.ops_import_services",
    PERSONNEL_UPLOAD: "app.admin.routes",
    HISTORY_IMPORT: "app.performance.history_import_routes",
}


def _xlsx(rows=(), edit=None) -> io.BytesIO:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(list(HEADERS))
    for row in rows:
        sheet.append(list(row))
    if edit is not None:
        edit(sheet)
    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer


def _far_cell(sheet):
    sheet["XFD1048576"] = "x"


def _bold_far_cell(sheet):
    sheet["H1048576"].font = Font(bold=True)


def _damaged() -> io.BytesIO:
    """A real .xlsx package (passes the upload validators) whose worksheet XML is broken."""
    source = zipfile.ZipFile(_xlsx())
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            broken = info.filename == "xl/worksheets/sheet1.xml"
            target.writestr(info.filename, b"<worksheet><sheetData><row>" if broken else source.read(info.filename))
    out.seek(0)
    return out


def _past_period(env) -> int:
    from app.extensions import db
    from app.models import PerformancePeriod

    with env.app.app_context():
        period = PerformancePeriod(title="Issue 104 geçmiş dönem", period_type="yillik", is_active=False,
                                   start_date=date(2024, 1, 1), end_date=date(2024, 12, 31))
        db.session.add(period)
        db.session.commit()
        return int(period.id)


def _post(env, client, path, upload, **form):
    data = {"excel_file": (upload, "personel.xlsx"), **form}
    if path == HISTORY_IMPORT:
        data.setdefault("period_id", str(_past_period(env)))
    return client.post(path, data=data, content_type="multipart/form-data", follow_redirects=True)


def _counts(env) -> dict[str, int]:
    from app.models import (
        AuditLog,
        OrganizationUnit,
        PerformanceImportBatch,
        PerformanceImportBatchRow,
        User,
    )

    with env.app.app_context():
        return {model.__name__: model.query.count()
                for model in (User, OrganizationUnit, AuditLog, PerformanceImportBatch, PerformanceImportBatchRow)}


@pytest.mark.parametrize("path", ROUTES)
@pytest.mark.parametrize(
    ("upload", "code"),
    [
        (lambda: _xlsx([_xlsx_row("R5104A", "personel")], edit=_far_cell), "EXCEL_TOO_MANY_ROWS"),
        (lambda: _xlsx(edit=lambda ws: ws.__setitem__("BM1", "fazla")), "EXCEL_TOO_MANY_COLUMNS"),
        (_damaged, "EXCEL_UNREADABLE"),
    ],
    ids=["far-cell", "far-column", "damaged"],
)
def test_hostile_workbook_is_refused_with_a_fixed_message_and_writes_nothing(env, path, upload, code):
    client = _web(env, "admin")
    before = _counts(env)
    response = _post(env, client, path, upload())
    assert response.status_code == 200
    assert MESSAGES[code] in response.get_data(as_text=True)
    assert _counts(env) == before
    assert _user_by_sicil(env, "R5104A") is None


@pytest.mark.parametrize("path", (ADMIN_IMPORT, PERSONNEL_UPLOAD))
def test_valid_rows_before_the_far_cell_are_not_written(env, path):
    """No partial import: the whole file is refused before the first database write."""
    client = _web(env, "admin")
    before = _counts(env)
    rows = [_xlsx_row("R5104B", "personel"), _xlsx_row("R5104C", "personel")]
    response = _post(env, client, path, _xlsx(rows, edit=lambda ws: ws.__setitem__("A20002", "fazla")),
                     activate_new_users="on")
    assert MESSAGES["EXCEL_TOO_MANY_ROWS"] in response.get_data(as_text=True)
    assert _counts(env) == before
    assert _user_by_sicil(env, "R5104B") is None and _user_by_sicil(env, "R5104C") is None


@pytest.mark.parametrize("path", (ADMIN_IMPORT, PERSONNEL_UPLOAD))
def test_valid_file_with_a_far_styled_cell_still_imports(env, path):
    client = _web(env, "admin")
    response = _post(env, client, path, _xlsx([_xlsx_row("R5104D", "personel")], edit=_bold_far_cell),
                     activate_new_users="on")
    assert response.status_code == 200
    created = _user_by_sicil(env, "R5104D")
    assert created is not None and created.role == "personel"


def test_valid_history_file_still_creates_a_preview_batch(env):
    from app.models import PerformanceImportBatch

    client = _web(env, "admin")
    sheet_rows = [("Sicil No", "Ad Soyad", "Toplam Puan"), (SICIL["plain"], "Kisi Plain", 88)]
    workbook = Workbook()
    for row in sheet_rows:
        workbook.active.append(list(row))
    workbook.active["H1048576"].font = Font(bold=True)
    upload = io.BytesIO()
    workbook.save(upload)
    upload.seek(0)
    response = _post(env, client, HISTORY_IMPORT, upload)
    assert response.status_code == 200
    with env.app.app_context():
        batches = PerformanceImportBatch.query.all()
        assert len(batches) == 1 and batches[0].row_count == 1


@pytest.mark.parametrize("path", ROUTES)
def test_non_admin_never_reaches_the_reader(env, path, monkeypatch):
    import importlib

    calls = []
    module = importlib.import_module(ROUTE_MODULES[path])
    real = module.read_import_rows

    def spy(upload):
        calls.append(upload)
        return real(upload)

    monkeypatch.setattr(module, "read_import_rows", spy)
    client = _web(env, "plain")
    before = _counts(env)
    response = client.post(path, data={"excel_file": (_xlsx(edit=_far_cell), "personel.xlsx"), "period_id": "1"},
                           content_type="multipart/form-data")
    assert response.status_code in {302, 403}
    assert calls == []
    after = _counts(env)
    assert after.pop("AuditLog") >= before.pop("AuditLog")  # the denial itself may be audited
    assert after == before
