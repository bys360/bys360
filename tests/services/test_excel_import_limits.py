"""Contract for app/services/excel_import_limits.py (issue #104: Excel import resource limits).

Before the fix the import routes called ``load_workbook(data_only=True)`` and
``list(ws.iter_rows())``: one value at XFD1048576 or a bold font on H1048576 in a 5 KB upload
made openpyxl build the whole rectangle (gigabytes, minutes). These tests craft such files in
memory and check that the reader refuses or bounds them with small memory use (tracemalloc,
so the budgets hold on Windows and Linux), and that valid workbooks read exactly as before.
"""
from __future__ import annotations

import inspect
import io
import logging
import struct
import tracemalloc
import zipfile
import zlib
from datetime import datetime
from types import SimpleNamespace
from typing import cast

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from app.services import excel_import_limits as limits
from app.services.excel_import_limits import ExcelImportRejected, read_import_rows

pytestmark = pytest.mark.ci_safe

NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
SHEET = "xl/worksheets/sheet1.xml"
MEMORY_BUDGET = 32 * 1024 * 1024  # the pre-fix code needed gigabytes for the hostile files below


# --------------------------------------------------------------------------- workbook builders


def _book(rows=(("Sicil No", "Ad"), ("1001", "Ayşe")), edit=None) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    for row in rows:
        sheet.append(list(row))
    if edit is not None:
        edit(sheet)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _patch(raw: bytes, replace: dict[str, bytes | None], add: dict[str, bytes] | None = None) -> bytes:
    """Copy a package, replacing (bytes) or dropping (None) members and adding new ones."""
    source = zipfile.ZipFile(io.BytesIO(raw))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            data = replace.get(info.filename, source.read(info.filename))
            if data is not None:
                target.writestr(info.filename, data)
        for name, data in (add or {}).items():
            target.writestr(name, data)
    return out.getvalue()


def _sheet_xml(sheet_data: str, *, head: str = "", tail: str = "") -> bytes:
    return (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            f'<worksheet xmlns="{NS}">{head}<sheetData>{sheet_data}</sheetData>{tail}</worksheet>').encode()


def _with_sheet(sheet_data: str, **kwargs) -> bytes:
    return _patch(_book(), {SHEET: _sheet_xml(sheet_data, **kwargs)})


def _inline(ref: str, text: str) -> str:
    return f'<c r="{ref}" t="inlineStr"><is><t>{text}</t></is></c>'


def _shared_strings_book(items: str, *, part: str = "xl/sharedStrings.xml") -> bytes:
    """A workbook whose A1 is shared string 0, with the table stored at ``part``."""
    raw = _book()
    types = zipfile.ZipFile(io.BytesIO(raw)).read("[Content_Types].xml").decode()
    override = (f'<Override PartName="/{part}" ContentType="application/vnd.openxmlformats-officedocument.'
                'spreadsheetml.sharedStrings+xml"/>')
    table = f'<?xml version="1.0" encoding="UTF-8"?><sst xmlns="{NS}">{items}</sst>'.encode()
    return _patch(raw, {"[Content_Types].xml": types.replace("</Types>", override + "</Types>").encode(),
                        SHEET: _sheet_xml('<row r="1"><c r="A1" t="s"><v>0</v></c></row>')},
                  add={part: table})


def _read(raw: bytes):
    return read_import_rows(io.BytesIO(raw))


def _rejected(raw: bytes) -> str:
    with pytest.raises(ExcelImportRejected) as caught:
        _read(raw)
    return caught.value.code


def _bounded(raw: bytes):
    """Read under tracemalloc; returns (rows or rejection code, peak bytes)."""
    tracemalloc.start()
    try:
        try:
            result: object = _read(raw)
        except ExcelImportRejected as exc:
            result = exc.code
        peak = tracemalloc.get_traced_memory()[1]
    finally:
        tracemalloc.stop()
    return result, peak


def _old_path_rows(raw: bytes):
    """What the import routes read before the fix (trailing empty rows/columns dropped)."""
    rows = list(load_workbook(io.BytesIO(raw), data_only=True).active.iter_rows(values_only=True))
    while rows and all(v is None for v in rows[-1]):
        rows.pop()
    width = max((i + 1 for row in rows for i, v in enumerate(row) if v is not None), default=0)
    return [tuple(row[:width]) for row in rows]


def _chunk(kind: bytes, data: bytes) -> bytes:
    return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data) & 0xFFFFFFFF)


def _png() -> bytes:
    raw = b"".join(b"\x00" + b"\x80" * 6 for _ in range(2))
    return (b"\x89PNG\r\n\x1a\n" + _chunk(b"IHDR", struct.pack(">IIBBBBB", 2, 2, 8, 2, 0, 0, 0))
            + _chunk(b"IDAT", zlib.compress(raw)) + _chunk(b"IEND", b""))


# --------------------------------------------------------------------------- far cells (#104)


def test_value_at_xfd1048576_is_refused_in_bounded_memory():
    result, peak = _bounded(_book(edit=lambda ws: ws.__setitem__("XFD1048576", "x")))
    assert result == "EXCEL_TOO_MANY_ROWS"
    assert peak < MEMORY_BUDGET


def test_bold_font_on_h1048576_does_not_inflate_the_read():
    def bold(ws):
        ws["H1048576"].font = Font(bold=True)

    result, peak = _bounded(_book(edit=bold))
    assert result == [("Sicil No", "Ad"), ("1001", "Ayşe")]
    assert peak < MEMORY_BUDGET


def test_row_limit_boundary():
    at_limit = _book(edit=lambda ws: ws.__setitem__(f"A{limits.MAX_SHEET_ROWS}", "son"))
    rows = _read(at_limit)
    assert len(rows) == limits.MAX_SHEET_ROWS == 20_001 and rows[-1] == ("son", None)
    over = _book(edit=lambda ws: ws.__setitem__(f"A{limits.MAX_SHEET_ROWS + 1}", "fazla"))
    assert _rejected(over) == "EXCEL_TOO_MANY_ROWS"


def test_column_limit_boundary():
    assert limits.MAX_COLUMNS == 64
    rows = _read(_book(edit=lambda ws: ws.__setitem__("BL1", "son")))  # column 64
    assert len(rows[0]) == 64 and rows[0][-1] == "son"
    assert _rejected(_book(edit=lambda ws: ws.__setitem__("BM1", "fazla"))) == "EXCEL_TOO_MANY_COLUMNS"


def test_forged_dimension_is_ignored():
    raw = _with_sheet('<row r="1">' + _inline("A1", "Sicil No") + '</row><row r="2">' + _inline("A2", "1001") + "</row>",
                      head='<dimension ref="A1:XFD1048576"/>')
    result, peak = _bounded(raw)
    assert result == [("Sicil No",), ("1001",)]
    assert peak < MEMORY_BUDGET


def test_sparse_data_in_a_wide_declared_range_reads_only_the_real_extent():
    raw = _with_sheet('<row r="1">' + _inline("A1", "a") + '</row><row r="2">' + _inline("B2", "b") + "</row>"
                      '<row r="20001"><c r="Z20001" s="1"/></row>', head='<dimension ref="A1:Z20001"/>')
    assert _read(raw) == [("a", None), (None, "b")]


def test_cells_without_a_value_do_not_extend_the_extent():
    """data_only reading: a formula without a cached value, an empty <v> or a styled cell is None."""
    raw = _with_sheet('<row r="1">' + _inline("A1", "a") + "</row>"
                      '<row r="30000"><c r="A30000"><f>1+1</f></c><c r="B30000"><v></v></c><c r="CZ30000" s="1"/></row>')
    assert _read(raw) == [("a",)]


def test_row_and_column_numbering_follows_openpyxl():
    raw = _with_sheet("<row>" + '<c t="inlineStr"><is><t>a</t></is></c><c t="inlineStr"><is><t>b</t></is></c>'
                      '</row><row r="3.0"><c r="C3"><v>7</v></c></row>')
    assert _read(raw) == [("a", "b", None), (None, None, None), (None, None, 7)]


def test_empty_sheet_reads_as_no_rows():
    assert _read(_book(rows=())) == []


# --------------------------------------------------------------------------- sheet structure


@pytest.mark.parametrize(
    ("sheet_data", "tail"),
    [
        ('<row r="3">' + '<c r="A3" s="1"/>' * (limits.MAX_ROW_CELLS + 1) + "</row>", ""),
        ('<row r="3"><c r="XFE3" s="1"/></row>', ""),
        ('<row r="1048577"><c r="A1048577" s="1"/></row>', ""),
        ('<row r="3"><c r="A3" t="inlineStr"><is>' + "<r><t>x</t></r>" * 600 + "</is></c></row>", ""),
        ("", '<mergeCells count="1">'
         + "".join(f'<mergeCell ref="A{n}:B{n}"/>' for n in range(3, limits.MAX_SHEET_EXTRA_ELEMENTS + 3)) + "</mergeCells>"),
    ],
    ids=["row-wider-than-excel", "column-past-xfd", "row-past-excel", "cell-with-too-many-elements", "merge-cell-flood"],
)
def test_sheet_structure_floods_are_refused_in_bounded_memory(sheet_data, tail):
    result, peak = _bounded(_with_sheet('<row r="1">' + _inline("A1", "a") + "</row>" + sheet_data, tail=tail))
    assert result == "EXCEL_SHEET_STRUCTURE_LIMIT"
    assert peak < MEMORY_BUDGET


def test_styled_empty_cells_and_rows_are_counted(monkeypatch):
    monkeypatch.setattr(limits, "MAX_SHEET_CELLS", 50)
    styled = "".join(f'<row r="{n}"><c r="A{n}" s="1"/></row>' for n in range(2, 60))
    assert _rejected(_with_sheet('<row r="1">' + _inline("A1", "a") + "</row>" + styled)) == "EXCEL_SHEET_STRUCTURE_LIMIT"
    monkeypatch.setattr(limits, "MAX_SHEET_CELLS", 10**9)
    monkeypatch.setattr(limits, "MAX_SHEET_ROW_ELEMENTS", 50)
    empty_rows = "".join(f'<row r="{n}"/>' for n in range(2, 60))
    assert _rejected(_with_sheet('<row r="1">' + _inline("A1", "a") + "</row>" + empty_rows)) == "EXCEL_SHEET_STRUCTURE_LIMIT"


# --------------------------------------------------------------------------- archive limits


def test_upload_over_10_mb_is_refused_before_unzipping():
    assert limits.MAX_UPLOAD_BYTES == 10 * 1024 * 1024
    assert _rejected(b"\0" * (limits.MAX_UPLOAD_BYTES + 1)) == "EXCEL_UPLOAD_TOO_LARGE"


def test_abnormal_compression_ratio_is_refused():
    padded = _sheet_xml('<row r="1">' + _inline("A1", "a") + "</row>").replace(b"<sheetData>", b"<!--" + b"a" * (2 * 1024 * 1024) + b"--><sheetData>")
    raw = _patch(_book(), {SHEET: padded})
    assert len(raw) < 64 * 1024
    result, peak = _bounded(raw)
    assert result == "EXCEL_ARCHIVE_LIMIT"
    assert peak < MEMORY_BUDGET


def test_too_many_archive_members_is_refused():
    raw = _patch(_book(), {}, add={f"xl/media/extra{n}.bin": b"x" for n in range(limits.MAX_ARCHIVE_MEMBERS)})
    assert _rejected(raw) == "EXCEL_ARCHIVE_LIMIT"


def _directory(infolist) -> zipfile.ZipFile:
    """Only the zip directory: check_archive reads nothing else."""
    return cast(zipfile.ZipFile, SimpleNamespace(infolist=infolist))


def test_archive_member_limit_boundary():
    tiny = SimpleNamespace(file_size=1, compress_size=1)
    limits.check_archive(_directory(lambda: [tiny] * limits.MAX_ARCHIVE_MEMBERS))
    with pytest.raises(ExcelImportRejected):
        limits.check_archive(_directory(lambda: [tiny] * (limits.MAX_ARCHIVE_MEMBERS + 1)))


def test_declared_uncompressed_total_is_checked():
    assert limits.MAX_UNCOMPRESSED_BYTES == 64 * 1024 * 1024
    info = SimpleNamespace(file_size=limits.MAX_UNCOMPRESSED_BYTES // 2 + 1, compress_size=limits.MAX_UNCOMPRESSED_BYTES)
    with pytest.raises(ExcelImportRejected):
        limits.check_archive(_directory(lambda: [info, info]))
    limits.check_archive(_directory(lambda: [info]))


def test_oversized_styles_part_is_refused_before_openpyxl_parses_it():
    raw = _book()
    styles = zipfile.ZipFile(io.BytesIO(raw)).read("xl/styles.xml")
    xfs = b"".join(b'<xf numFmtId="%d" fontId="0" fillId="0" borderId="0" xfId="0"/>' % (n % 160) for n in range(70_000))
    bomb = styles.replace(b"</styleSheet>", b"<extra>" + xfs + b"</extra></styleSheet>")
    assert len(bomb) > limits.MAX_PART_BYTES
    result, peak = _bounded(_patch(raw, {"xl/styles.xml": bomb}))
    assert result == "EXCEL_ARCHIVE_LIMIT"
    assert peak < MEMORY_BUDGET


def test_oversized_workbook_part_is_refused_even_under_a_worksheet_name():
    """The workbook part is found through the manifest; renaming it does not lift the cap."""
    raw = _book()
    package = zipfile.ZipFile(io.BytesIO(raw))
    names = b"".join(b'<definedName name="n%d">Sheet!$A$1</definedName>' % n for n in range(120_000))
    workbook = package.read("xl/workbook.xml").replace(b"</workbook>", b"<definedNames>" + names + b"</definedNames></workbook>")
    types = package.read("[Content_Types].xml").replace(b"/xl/workbook.xml", b"/xl/worksheets/sheet9.xml")
    rels = package.read("_rels/.rels").replace(b"xl/workbook.xml", b"xl/worksheets/sheet9.xml")
    sheet_rels = package.read("xl/_rels/workbook.xml.rels").replace(b'Target="', b'Target="/xl/')
    hostile = _patch(raw, {"xl/workbook.xml": None, "[Content_Types].xml": types, "_rels/.rels": rels,
                           "xl/_rels/workbook.xml.rels": None},
                     add={"xl/worksheets/sheet9.xml": workbook, "xl/worksheets/_rels/sheet9.xml.rels": sheet_rels})
    result, peak = _bounded(hostile)
    assert result == "EXCEL_ARCHIVE_LIMIT"
    assert peak < MEMORY_BUDGET


def test_shared_strings_are_counted_at_the_manifest_path(monkeypatch):
    raw = _shared_strings_book("<si><t>ad</t></si>" * 40, part="xl/worksheets/sheet9.xml")
    assert _read(raw) == [("ad",)]
    monkeypatch.setattr(limits, "MAX_SHARED_STRINGS", 39)
    assert _rejected(raw) == "EXCEL_ARCHIVE_LIMIT"


@pytest.mark.parametrize(
    "items",
    ["<si><t>a</t></si>" + "<pad/>" * (limits.MAX_SHARED_STRING_EXTRA_ELEMENTS + 1),
     "<si>" + "<r><t>a</t></r>" * 600 + "</si>"],
    ids=["padding-outside-items", "item-with-too-many-runs"],
)
def test_padded_shared_string_table_is_refused(items):
    assert _rejected(_shared_strings_book(items)) == "EXCEL_ARCHIVE_LIMIT"


# --------------------------------------------------------------------------- damaged files


def _billion_laughs(root: str) -> bytes:
    return (b'<?xml version="1.0"?><!DOCTYPE x [<!ENTITY a "aaaaaaaaaa"><!ENTITY b "&a;&a;&a;&a;&a;&a;&a;&a;">]>'
            + f'<{root} xmlns="{NS}">&b;</{root}>'.encode())


def _broken_sheet_rels(raw: bytes) -> bytes:
    rels = zipfile.ZipFile(io.BytesIO(raw)).read("xl/_rels/workbook.xml.rels")
    return _patch(raw, {"xl/_rels/workbook.xml.rels": rels.replace(b"worksheets/sheet1.xml", b"worksheets/missing.xml")})


@pytest.mark.parametrize(
    "raw",
    [
        b"not an excel file",
        _book()[:500],
        _patch(_book(), {"xl/workbook.xml": None}),
        _patch(_book(), {SHEET: b"<worksheet><sheetData><row>"}),
        _patch(_book(), {SHEET: _billion_laughs("worksheet")}),
        _patch(_book(), {"xl/styles.xml": _billion_laughs("styleSheet")}),
        _with_sheet('<row r="1"><c r="A1" t="s"><v>99</v></c></row>'),
        _with_sheet('<row r="1"><c r="A$" t="inlineStr"><is><t>a</t></is></c></row>'),
        _with_sheet('<row r="x1"><c t="inlineStr"><is><t>a</t></is></c></row>'),
        _broken_sheet_rels(_book()),
    ],
    ids=["not-zip", "truncated-zip", "no-workbook-part", "broken-sheet-xml", "sheet-doctype", "styles-doctype",
         "shared-string-out-of-range", "bad-cell-ref", "bad-row-number", "missing-sheet-part"],
)
def test_damaged_workbook_gives_the_fixed_message(raw):
    with pytest.raises(ExcelImportRejected) as caught:
        _read(raw)
    assert caught.value.code == "EXCEL_UNREADABLE"
    assert caught.value.user_message == limits.MESSAGES["EXCEL_UNREADABLE"]


@pytest.mark.parametrize("part", [SHEET, "xl/styles.xml", "xl/workbook.xml"])
def test_doctype_is_refused_by_the_reader_itself(part, caplog):
    """Production has no defusedxml: the DOCTYPE probe, not the XML parser, must refuse it."""
    raw = _book()
    original = zipfile.ZipFile(io.BytesIO(raw)).read(part)
    assert not original.startswith(b"<?xml")
    with_doctype = b'<!DOCTYPE x [<!ENTITY a "a">]>' + original
    caplog.set_level(logging.WARNING, logger=limits.__name__)
    assert _rejected(_patch(raw, {part: with_doctype})) == "EXCEL_UNREADABLE"
    assert caplog.records[-1].getMessage().endswith("code=EXCEL_UNREADABLE")  # no parser exception involved


@pytest.mark.parametrize("part", [SHEET, "xl/styles.xml"])
@pytest.mark.parametrize("encoding", ["utf-8", "utf-16"])
def test_doctype_far_into_the_part_or_in_utf16_is_refused(part, encoding, caplog):
    """A long leading comment or a UTF-16 part must not move the DOCTYPE out of the probe's sight."""
    raw = _book()
    body = zipfile.ZipFile(io.BytesIO(raw)).read(part).decode()
    text = f'<?xml version="1.0" encoding="{encoding}"?><!--{"x" * 70_000}--><!DOCTYPE x [<!ENTITY a "a">]>{body}'
    caplog.set_level(logging.WARNING, logger=limits.__name__)
    assert _rejected(_patch(raw, {part: text.encode(encoding)})) == "EXCEL_UNREADABLE"
    assert caplog.records[-1].getMessage().endswith("code=EXCEL_UNREADABLE")


def test_utf16_parts_without_doctype_still_read():
    raw = _book()
    body = zipfile.ZipFile(io.BytesIO(raw)).read(SHEET).decode()
    utf16 = ('<?xml version="1.0" encoding="utf-16"?>' + body).encode("utf-16")
    assert _read(_patch(raw, {SHEET: utf16})) == [("Sicil No", "Ad"), ("1001", "Ayşe")]


def _two_sheets(second: bytes) -> bytes:
    def add_sheet(ws):
        ws.parent.create_sheet("Açıklama")["A1"] = "yer tutucu"

    return _patch(_book(edit=add_sheet), {"xl/worksheets/sheet2.xml": second})


def _varied_rows(count: int) -> str:
    return "".join(f'<row r="{n}"><c r="A{n}"><v>{n * 7919 % 100003}</v></c></row>' for n in range(1, count))


def test_large_inactive_sheet_with_a_dimension_loads():
    """openpyxl opens every sheet while loading only to find <dimension>; that read stays short."""
    second = _sheet_xml(_varied_rows(120_000), head='<dimension ref="A1:A119999"/>')
    assert len(second) > limits.MAX_PART_BYTES
    assert _read(_two_sheets(second)) == [("Sicil No", "Ad"), ("1001", "Ayşe")]


def test_large_sheet_without_a_dimension_is_refused_while_loading():
    """Without <dimension> openpyxl would parse the whole sheet at load time; the part budget stops it."""
    second = _sheet_xml(_varied_rows(120_000))
    assert len(second) > limits.MAX_PART_BYTES
    result, peak = _bounded(_two_sheets(second))
    assert result == "EXCEL_ARCHIVE_LIMIT"
    assert peak < MEMORY_BUDGET


def test_rejections_log_a_code_and_never_file_content(caplog):
    caplog.set_level(logging.WARNING, logger=limits.__name__)
    secret = "GIZLI-HUCRE-ICERIGI"
    with pytest.raises(ExcelImportRejected):
        _read(_with_sheet('<row r="1">' + _inline("A1", secret) + '</row><row r="90000"><c r="A90000"><v>1</v></c></row>'))
    with pytest.raises(ExcelImportRejected):
        _read(_patch(_book(rows=((secret,),)), {SHEET: b"<worksheet>" + secret.encode()}))
    text = caplog.text
    assert "code=EXCEL_TOO_MANY_ROWS" in text and "code=EXCEL_UNREADABLE" in text and "exc_type=" in text
    assert secret not in text


def test_messages_are_turkish_and_say_nothing_was_written():
    for message in limits.MESSAGES.values():
        assert message.endswith("Hiçbir kayıt veritabanına yazılmadı.")
    assert "20.000" in limits.MESSAGES["EXCEL_TOO_MANY_ROWS"]


# --------------------------------------------------------------------------- valid workbooks


def test_excel_style_workbook_reads_exactly_as_before():
    """XlsxWriter writes like Excel: shared strings, numbers, dates, booleans, cached formulas."""
    import xlsxwriter

    buffer = io.BytesIO()
    book = xlsxwriter.Workbook(buffer, {"in_memory": True})
    sheet = book.add_worksheet("Personel")
    date_format = book.add_format({"num_format": "dd.mm.yyyy"})
    sheet.write_row(0, 0, ["Sicil No", "Ad", "Soyad", "Puan", "Aktif", "Tarih", "Toplam"])
    sheet.write_row(1, 0, ["1001", "Ayşe", "Yılmaz", 87.5, True])
    sheet.write_datetime(1, 5, datetime(2026, 1, 15), date_format)
    sheet.write_formula(1, 6, "=D2*2", None, 175)
    sheet.write_row(3, 0, ["1002", "Çağrı", "Öztürk", 90, False])
    sheet.set_column(0, 30, 18, book.add_format({"bold": True}))
    book.close()
    raw = buffer.getvalue()
    assert "xl/sharedStrings.xml" in zipfile.ZipFile(io.BytesIO(raw)).namelist()
    rows = _read(raw)
    assert rows == _old_path_rows(raw)
    assert rows[1] == ("1001", "Ayşe", "Yılmaz", 87.5, True, datetime(2026, 1, 15), 175)
    assert rows[2] == (None,) * 7


def test_workbook_with_embedded_pictures_reads():
    from openpyxl.drawing.image import Image as XImage
    from PIL import Image

    jpeg = io.BytesIO()
    Image.new("RGB", (4, 4), (10, 120, 200)).save(jpeg, format="JPEG")

    def pictures(ws):
        ws.add_image(XImage(io.BytesIO(_png())), "D2")
        ws.add_image(XImage(io.BytesIO(jpeg.getvalue())), "G2")

    raw = _book(edit=pictures)
    assert [n for n in zipfile.ZipFile(io.BytesIO(raw)).namelist() if n.startswith("xl/media/")]
    assert _read(raw) == [("Sicil No", "Ad"), ("1001", "Ayşe")]


def test_official_personnel_template_reads_like_before():
    from app.services.personnel.excel_template import (
        PERSONNEL_EXCEL_TEMPLATE_HEADERS,
        build_personnel_import_template_workbook,
    )

    buffer = io.BytesIO()
    build_personnel_import_template_workbook().save(buffer)
    raw = buffer.getvalue()
    rows = _read(raw)
    assert rows[0] == PERSONNEL_EXCEL_TEMPLATE_HEADERS
    assert rows == _old_path_rows(raw)  # the 50 pre-formatted empty rows are dropped, data rows kept


def test_valid_file_at_the_row_limit_with_data_in_every_row():
    rows = [("Sicil No", "Ad", "Puan")] + [(f"S{n:05d}", f"Kişi {n}", n % 100) for n in range(limits.MAX_DATA_ROWS)]
    raw = _book(rows=rows)
    assert _read(raw) == [tuple(r) for r in rows]


def test_file_storage_stream_is_accepted():
    from werkzeug.datastructures import FileStorage

    upload = FileStorage(stream=io.BytesIO(_book()), filename="personel.xlsx")
    assert read_import_rows(upload) == [("Sicil No", "Ad"), ("1001", "Ayşe")]


# --------------------------------------------------------------------------- openpyxl coupling


def test_openpyxl_internals_used_by_the_reader_are_present():
    """Pinned openpyxl 3.1.5: the reader relies on these; an upgrade that drops one fails here."""
    from openpyxl.reader.excel import ExcelReader
    from openpyxl.worksheet._read_only import ReadOnlyWorksheet

    assert list(inspect.signature(ExcelReader.__init__).parameters)[:7] == [
        "self", "fn", "read_only", "keep_vba", "data_only", "keep_links", "rich_text"]
    for name in ("read_manifest", "read"):
        assert callable(getattr(ExcelReader, name))
    assert "_worksheet_path" in inspect.getsource(ReadOnlyWorksheet.__init__)
    assert "self.parent._archive.open(self._worksheet_path)" in inspect.getsource(ReadOnlyWorksheet._get_source)
