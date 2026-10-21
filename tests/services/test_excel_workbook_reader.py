"""Unit contract for app.services.excel_workbook_reader (no app, no database)."""
from __future__ import annotations

import io
import logging
import zipfile
from xml.etree.ElementTree import ParseError

import pytest
from PIL.Image import DecompressionBombError

from app.services import excel_workbook_reader as reader


def _valid() -> bytes:
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.active.append(["sicil_no", "ad"])
    workbook.active.append(["T1", "Deneme"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


@pytest.fixture(autouse=True)
def _logging_enabled():
    previous = logging.root.manager.disable
    logging.disable(logging.NOTSET)
    yield
    logging.disable(previous)


def test_valid_workbook_returns_value_rows():
    assert reader.read_active_sheet_rows(io.BytesIO(_valid())) == [("sicil_no", "ad"), ("T1", "Deneme")]


@pytest.mark.parametrize(
    "payload",
    [b"", b"bu bir excel dosyasi degil" * 10, bytes.fromhex("D0CF11E0A1B11AE1") + b"\x00" * 504],
    ids=["empty", "text", "ole-encrypted"],
)
def test_non_zip_bytes_raise_the_domain_error_without_chain(payload):
    with pytest.raises(reader.ExcelWorkbookReadError) as caught:
        reader.read_active_sheet_rows(io.BytesIO(payload))
    assert caught.value.__cause__ is None and caught.value.__suppress_context__
    assert caught.value.user_message == reader.EXCEL_WORKBOOK_UNREADABLE_MESSAGE


@pytest.mark.parametrize(
    "error",
    [
        ValueError("SENTINEL"), KeyError("SENTINEL"), TypeError("SENTINEL"), OSError("SENTINEL"),
        IndexError("SENTINEL"), RuntimeError("SENTINEL"), EOFError("SENTINEL"), ParseError("SENTINEL"),
        zipfile.BadZipFile("SENTINEL"), DecompressionBombError("SENTINEL"),
    ],
    ids=lambda e: type(e).__name__,
)
def test_parser_errors_are_converted_and_logged_by_class_only(monkeypatch, caplog, error):
    def boom(*_args, **_kwargs):
        raise error

    monkeypatch.setattr(reader, "load_workbook", boom)
    caplog.set_level(logging.DEBUG)
    with pytest.raises(reader.ExcelWorkbookReadError):
        reader.read_active_sheet_rows(io.BytesIO(_valid()))
    assert "EXCEL_WORKBOOK_UNREADABLE" in caplog.text
    assert type(error).__qualname__ in caplog.text
    assert "SENTINEL" not in caplog.text
    assert not any(record.exc_info for record in caplog.records)


@pytest.mark.parametrize("error", [AttributeError("bug"), NameError("bug"), ImportError("bug"), MemoryError()],
                         ids=lambda e: type(e).__name__)
def test_programming_errors_are_not_swallowed(monkeypatch, error):
    def boom(*_args, **_kwargs):
        raise error

    monkeypatch.setattr(reader, "load_workbook", boom)
    with pytest.raises(type(error)):
        reader.read_active_sheet_rows(io.BytesIO(_valid()))


def test_workbook_without_a_worksheet_is_unreadable(monkeypatch, caplog):
    class _NoActive:
        active = None

    monkeypatch.setattr(reader, "load_workbook", lambda *_a, **_k: _NoActive())
    caplog.set_level(logging.DEBUG)
    with pytest.raises(reader.ExcelWorkbookReadError):
        reader.read_active_sheet_rows(io.BytesIO(b"ignored"))
    assert "EXCEL_WORKBOOK_NO_WORKSHEET" in caplog.text
