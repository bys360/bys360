"""The pinned Pillow works with the libraries that hand it pictures: openpyxl, XlsxWriter, ReportLab.

The application does not process images with PIL itself (it only imports
``DecompressionBombError`` to refuse hostile uploads); Pillow is reached through openpyxl
(pictures inside an uploaded or generated .xlsx) and ReportLab (pictures in a PDF). XlsxWriter
parses PNG headers on its own. A Pillow major upgrade (12.3.0 -> 13.0.0 removed ``Image._show``, moved
``get_child_images``, stopped clipping I-mode PNG saves) must keep these paths working, and the
suite must really run on the pinned version.
"""
from __future__ import annotations

import io
import re
import zipfile
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parents[2]


def _pinned(manifest: str, name: str) -> str:
    text = (REPO_ROOT / manifest).read_text(encoding="utf-8")
    match = re.search(rf"^{name}==([^\s\\]+)", text, flags=re.IGNORECASE | re.MULTILINE)
    assert match, f"{name} is not pinned in {manifest}"
    return match.group(1)


def _png(mode="RGB", size=(24, 16)) -> bytes:
    from PIL import Image

    color = (10, 120, 200, 128) if mode == "RGBA" else (10, 120, 200)
    buffer = io.BytesIO()
    Image.new(mode, size, color).save(buffer, "PNG")
    return buffer.getvalue()


def test_the_suite_runs_on_the_pillow_version_both_manifests_pin():
    import PIL

    assert _pinned("requirements.txt", "Pillow") == _pinned("requirements.lock", "pillow") == PIL.__version__


def test_openpyxl_writes_and_reads_back_a_png_picture():
    from openpyxl import Workbook, load_workbook
    from openpyxl.drawing.image import Image as XLImage
    from PIL import Image

    workbook = Workbook()
    workbook.active.append(["Sicil No", "Ad"])
    workbook.active.add_image(XLImage(io.BytesIO(_png())), "D2")
    buffer = io.BytesIO()
    workbook.save(buffer)

    media = [name for name in zipfile.ZipFile(io.BytesIO(buffer.getvalue())).namelist() if name.startswith("xl/media/")]
    assert media == ["xl/media/image1.png"]
    pictures = load_workbook(io.BytesIO(buffer.getvalue())).active._images
    assert len(pictures) == 1
    reread = Image.open(io.BytesIO(pictures[0]._data()))
    assert (reread.format, reread.size) == ("PNG", (24, 16))


def test_openpyxl_drops_an_unreadable_picture_but_keeps_the_cells():
    import warnings

    from openpyxl import Workbook, load_workbook
    from openpyxl.drawing.image import Image as XLImage

    workbook = Workbook()
    workbook.active.append(["Sicil No", "Ad"])
    workbook.active.add_image(XLImage(io.BytesIO(_png())), "D2")
    buffer = io.BytesIO()
    workbook.save(buffer)
    source = zipfile.ZipFile(io.BytesIO(buffer.getvalue()))
    damaged = io.BytesIO()
    with zipfile.ZipFile(damaged, "w") as target:
        for info in source.infolist():
            data = source.read(info.filename)
            target.writestr(info.filename, b"\x89PNG\r\n\x1a\nnot a picture" if info.filename.startswith("xl/media/") else data)

    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter("always")
        sheet = load_workbook(io.BytesIO(damaged.getvalue())).active
    assert sheet._images == []
    assert [row for row in sheet.iter_rows(values_only=True)] == [("Sicil No", "Ad")]
    assert any("cannot be read" in str(item.message) for item in caught)


def test_xlsxwriter_embeds_a_png_picture():
    import xlsxwriter

    picture = _png()
    buffer = io.BytesIO()
    workbook = xlsxwriter.Workbook(buffer, {"in_memory": True})
    workbook.add_worksheet().insert_image("B2", "logo.png", {"image_data": io.BytesIO(picture)})
    workbook.close()

    archive = zipfile.ZipFile(io.BytesIO(buffer.getvalue()))
    assert archive.read("xl/media/image1.png") == picture


def test_reportlab_draws_rgb_and_transparent_png_pictures():
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas
    from reportlab.platypus import Image as RLImage, Paragraph, SimpleDocTemplate

    page = io.BytesIO()
    pdf = canvas.Canvas(page)
    pdf.drawImage(ImageReader(io.BytesIO(_png())), 20, 20, 24, 16)
    pdf.drawImage(ImageReader(io.BytesIO(_png("RGBA"))), 60, 20, 24, 16, mask="auto")
    pdf.save()
    assert page.getvalue().startswith(b"%PDF")
    assert page.getvalue().count(b"/Subtype /Image") >= 2

    document = io.BytesIO()
    SimpleDocTemplate(document).build(
        [Paragraph("BYS360", getSampleStyleSheet()["Normal"]), RLImage(io.BytesIO(_png()), 24, 16)]
    )
    assert document.getvalue().startswith(b"%PDF")
    assert b"/Subtype /Image" in document.getvalue()
