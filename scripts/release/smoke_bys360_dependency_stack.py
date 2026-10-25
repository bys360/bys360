"""Smoke-test an environment installed from requirements.lock (run it with that environment's python).

Checks, in order, and stops at the first failure with a non-zero exit code:
1. every package pinned in requirements.lock is installed at exactly the pinned version;
2. the runtime modules import;
3. Pillow works through openpyxl, XlsxWriter and ReportLab (PNG in .xlsx and in PDF), and
   pandas reads an .xlsx through openpyxl;
4. the application factory starts on an in-memory SQLite database and serves /login.

No network, no database server, no files outside a temporary directory.
"""
from __future__ import annotations

import argparse
import importlib
import io
import json
import os
import re
import sys
import zipfile
from importlib import metadata
from pathlib import Path

PACKAGE = "BYS360_DEPENDENCY_STACK_SMOKE_V1"
LOCK_ENTRY_RE = re.compile(r"^([A-Za-z0-9][A-Za-z0-9._-]*)==([A-Za-z0-9][A-Za-z0-9._!+-]*)\s*\\?$")
RUNTIME_MODULES = (
    "flask", "flask_login", "flask_sqlalchemy", "flask_wtf", "flask_migrate", "flask_limiter", "sqlalchemy",
    "psycopg2", "dotenv", "wtforms", "email_validator", "openpyxl", "pandas", "waitress", "xlsxwriter", "PIL",
    "cryptography", "reportlab", "sentry_sdk", "redis", "prometheus_client", "werkzeug", "urllib3", "certifi",
)


def canonical_name(name: str) -> str:
    return re.sub(r"[-_.]+", "-", name).lower()


def check_pins(lock: Path) -> dict[str, str]:
    pins = {}
    for raw in lock.read_text(encoding="utf-8").splitlines():
        match = LOCK_ENTRY_RE.match(raw.strip())
        if match:
            pins[canonical_name(match.group(1))] = match.group(2)
    if not pins:
        raise AssertionError(f"{lock} has no pinned entries")
    wrong = {}
    for name, version in pins.items():
        try:
            installed = metadata.version(name)
        except metadata.PackageNotFoundError:
            installed = None
        if installed != version:
            wrong[name] = (version, installed)
    if wrong:
        raise AssertionError(f"installed versions differ from requirements.lock (pinned, installed): {wrong}")
    return pins


def check_imports() -> None:
    for module in RUNTIME_MODULES:
        importlib.import_module(module)


def _png(mode: str = "RGB") -> bytes:
    from PIL import Image

    buffer = io.BytesIO()
    Image.new(mode, (24, 16), (10, 120, 200, 128) if mode == "RGBA" else (10, 120, 200)).save(buffer, "PNG")
    return buffer.getvalue()


def check_image_stack() -> None:
    import pandas as pd
    import xlsxwriter
    from openpyxl import Workbook, load_workbook
    from openpyxl.drawing.image import Image as XLImage
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas

    workbook = Workbook()
    workbook.active.append(["Sicil No", "Ad"])
    workbook.active.append(["1001", "Ayse"])
    workbook.active.add_image(XLImage(io.BytesIO(_png())), "D2")
    xlsx = io.BytesIO()
    workbook.save(xlsx)
    assert len(load_workbook(io.BytesIO(xlsx.getvalue())).active._images) == 1
    frame = pd.read_excel(io.BytesIO(xlsx.getvalue()), engine="openpyxl", dtype=str)
    assert list(frame.columns) == ["Sicil No", "Ad"] and frame.iloc[0].tolist() == ["1001", "Ayse"]

    written = io.BytesIO()
    book = xlsxwriter.Workbook(written, {"in_memory": True})
    book.add_worksheet().insert_image("B2", "logo.png", {"image_data": io.BytesIO(_png())})
    book.close()
    assert "xl/media/image1.png" in zipfile.ZipFile(io.BytesIO(written.getvalue())).namelist()

    pdf_bytes = io.BytesIO()
    pdf = canvas.Canvas(pdf_bytes)
    pdf.drawImage(ImageReader(io.BytesIO(_png())), 20, 20, 24, 16)
    pdf.drawImage(ImageReader(io.BytesIO(_png("RGBA"))), 60, 20, 24, 16, mask="auto")
    pdf.save()
    assert pdf_bytes.getvalue().startswith(b"%PDF") and pdf_bytes.getvalue().count(b"/Subtype /Image") >= 2


def check_app(root: Path) -> None:
    os.environ.update(
        {
            "APP_ENV": "testing",
            "FLASK_ENV": "testing",
            "DATABASE_URL": "sqlite:///:memory:",
            "SECRET_KEY": "dependency-smoke-test-secret-key",
            "DEFAULT_FIRST_LOGIN_PASSWORD": "dependency-smoke-test-first-login",
            "FLASK_SKIP_SCHEMA_VALIDATION": "1",
            "AUTO_REPAIR_SCHEMA": "false",
            "STRICT_SCHEMA_CHECK": "false",
            "REQUIRE_DOTENV_FILE": "false",
            "STRICT_ENV_VALIDATION": "false",
            "SCHEDULER_ENABLED": "false",
            "MAIL_SUPPRESS_SEND": "true",
        }
    )
    sys.path.insert(0, str(root))
    from app import create_app

    app = create_app()
    app.config.update(TESTING=True)
    response = app.test_client().get("/login")
    assert response.status_code == 200, f"/login returned {response.status_code}"


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.split("\n", 1)[0])
    parser.add_argument("--root", default=".", help="Depo koku (varsayilan: mevcut dizin)")
    parser.add_argument("--lock", default="requirements.lock")
    args = parser.parse_args(argv)
    root = Path(args.root).resolve()
    steps = (
        ("pins", lambda: check_pins(root / args.lock)),
        ("imports", check_imports),
        ("image_stack", check_image_stack),
        ("app", lambda: check_app(root)),
    )
    result: dict[str, object] = {"package": PACKAGE, "python": sys.version.split()[0], "platform": sys.platform}
    for name, step in steps:
        try:
            value = step()
        except Exception as exc:  # report which step failed, then fail the job
            result.update({"ok": False, "failed_step": name, "error": f"{type(exc).__name__}: {exc}"})
            print(json.dumps(result, ensure_ascii=False, indent=2))
            return 1
        if name == "pins" and isinstance(value, dict):
            result["pinned_packages"] = len(value)
    import PIL

    result.update({"ok": True, "pillow": PIL.__version__})
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
